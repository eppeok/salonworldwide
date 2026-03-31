import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import datetime, timedelta
import os

wb = openpyxl.Workbook()

# Styles
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_font = Font(name='Calibri', size=10)
cell_align = Alignment(vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC')
)
green_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
red_fill = PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid')
yellow_fill = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')
total_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')

def style_header(ws, row=1, max_col=None):
    if max_col is None:
        max_col = ws.max_column
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

def style_cells(ws, start_row=2, end_row=None, max_col=None):
    if end_row is None:
        end_row = ws.max_row
    if max_col is None:
        max_col = ws.max_column
    for row in range(start_row, end_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = thin_border

# ===================== SHEET 1: LEAD DATABASE =====================
ws1 = wb.active
ws1.title = "Lead Database"

headers1 = ["Lead ID", "Full Name", "Salon Name", "City", "State", "Platform Found",
            "Instagram Handle", "Email", "Phone", "Website", "Follower Count",
            "Running Ads?", "Lead Score", "Lead Tier", "Current Phase",
            "First Contact Date", "First Contact Channel", "First Response Date",
            "Response Status", "Follow-Up Count", "Last Follow-Up Date",
            "Landing Page Click?", "Purchased?", "Purchase Date", "Notes"]

for i, h in enumerate(headers1, 1):
    ws1.cell(row=1, column=i, value=h)

# Sample data
samples = [
    ["L001", "Maria Rodriguez", "Glow Beauty Salon", "Miami", "FL", "Instagram", "@glowbeautymiami", "maria@glowbeauty.com", "305-555-0123", "glowbeautysalon.com", 3200, "Yes", 75, "HOT", "Attention", "2026-04-01", "Instagram DM", "2026-04-02", "Replied - Interested", 2, "2026-04-05", "Yes", "No", "", "Very interested, asked about reactivation module"],
    ["L002", "Jennifer Kim", "JK Hair Studio", "Los Angeles", "CA", "Apollo.io", "@jkhairstudio", "jen@jkhairstudio.com", "310-555-0456", "jkhairstudio.com", 5600, "Yes", 55, "WARM", "Interest", "2026-04-01", "Cold Email", "2026-04-04", "Replied - Maybe", 1, "2026-04-04", "No", "No", "", "Opened email 1 and 2, replied asking for more info"],
    ["L003", "Tasha Williams", "Crown & Glory", "Atlanta", "GA", "Facebook Group", "@crownandgloryatl", "tasha@crownandglory.com", "404-555-0789", "crownandglorysalon.com", 1800, "Unknown", 35, "COOL", "Attention", "2026-04-03", "Facebook DM", "", "No Response", 1, "2026-04-07", "No", "No", "", "Engaged with FB group post but no DM reply"],
    ["L004", "Fatima Hassan", "Luxe Locks Salon", "Houston", "TX", "Google Maps", "@luxelockshtx", "fatima@luxelocks.com", "713-555-0321", "luxelockssalon.com", 7800, "Yes", 85, "HOT", "Desire", "2026-04-01", "Instagram DM", "2026-04-01", "Clicked Link", 3, "2026-04-10", "Yes", "No", "", "Clicked link twice, sent WhatsApp voice note"],
    ["L005", "David Chen", "Studio 88 Hair", "San Francisco", "CA", "LinkedIn", "", "david@studio88hair.com", "415-555-0654", "studio88hair.com", 2100, "No", 25, "COOL", "Attention", "2026-04-05", "LinkedIn", "", "No Response", 0, "", "No", "No", "", "Connected on LinkedIn, no reply to message"],
]

for i, row in enumerate(samples, 2):
    for j, val in enumerate(row, 1):
        ws1.cell(row=i, column=j, value=val)

style_header(ws1, max_col=25)
style_cells(ws1, end_row=6, max_col=25)

# Column widths
widths1 = [8, 18, 20, 14, 8, 14, 18, 25, 15, 22, 12, 12, 10, 10, 12, 14, 16, 16, 20, 12, 14, 14, 10, 12, 30]
for i, w in enumerate(widths1, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

# Data validations
dv_platform = DataValidation(type="list", formula1='"Instagram,Facebook,LinkedIn,Google Maps,Yelp,Apollo.io,Hunter.io,Referral,Threads,Reddit"')
dv_ads = DataValidation(type="list", formula1='"Yes,No,Unknown"')
dv_tier = DataValidation(type="list", formula1='"HOT,WARM,COOL,COLD"')
dv_phase = DataValidation(type="list", formula1='"Attention,Interest,Desire,Action"')
dv_status = DataValidation(type="list", formula1='"No Response,Replied - Interested,Replied - Not Interested,Replied - Maybe,Clicked Link,Purchased"')
dv_yn = DataValidation(type="list", formula1='"Yes,No"')

ws1.add_data_validation(dv_platform)
ws1.add_data_validation(dv_ads)
ws1.add_data_validation(dv_tier)
ws1.add_data_validation(dv_phase)
ws1.add_data_validation(dv_status)
ws1.add_data_validation(dv_yn)

dv_platform.add(f"F2:F500")
dv_ads.add(f"L2:L500")
dv_tier.add(f"N2:N500")
dv_phase.add(f"O2:O500")
dv_status.add(f"S2:S500")
dv_yn.add(f"V2:V500")
dv_yn.add(f"W2:W500")

ws1.freeze_panes = 'A2'
ws1.auto_filter.ref = f"A1:Y{ws1.max_row}"

# ===================== SHEET 2: DAILY ACTIVITY LOG =====================
ws2 = wb.create_sheet("Daily Activity Log")

headers2 = ["Date", "Day", "Week #", "Phase", "IG DMs Sent", "IG DM Replies",
            "DM Reply Rate", "Emails Sent (GHL)", "Email Opens", "Email Open Rate",
            "Email Replies", "Email Click-Throughs", "LinkedIn Requests", "LinkedIn Accepted",
            "LinkedIn Messages", "LinkedIn Replies", "FB Group Comments", "FB Group Posts",
            "WhatsApp Sent", "WhatsApp Replies", "Landing Page Visits", "Sales",
            "Revenue", "Notes"]

for i, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=i, value=h)

# Pre-fill 42 days
start_date = datetime(2026, 4, 6)  # A Monday
phases = {1: "Attention", 2: "Attention", 3: "Interest", 4: "Interest", 5: "Desire", 6: "Action"}
days_of_week = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

for day_num in range(42):
    row = day_num + 2
    d = start_date + timedelta(days=day_num)
    week_num = (day_num // 7) + 1
    
    ws2.cell(row=row, column=1, value=d.strftime("%Y-%m-%d"))
    ws2.cell(row=row, column=2, value=days_of_week[day_num % 7])
    ws2.cell(row=row, column=3, value=week_num)
    ws2.cell(row=row, column=4, value=phases.get(week_num, "Action"))
    
    # Formula columns
    ws2.cell(row=row, column=7).value = f'=IF(E{row}=0,"",F{row}/E{row})'
    ws2.cell(row=row, column=7).number_format = '0.0%'
    ws2.cell(row=row, column=10).value = f'=IF(H{row}=0,"",I{row}/H{row})'
    ws2.cell(row=row, column=10).number_format = '0.0%'
    ws2.cell(row=row, column=23).value = f'=V{row}*97'
    ws2.cell(row=row, column=23).number_format = '$#,##0'

# Totals row
total_row = 44
ws2.cell(row=total_row, column=1, value="TOTALS")
ws2.cell(row=total_row, column=1).font = Font(name='Calibri', bold=True, size=11)
for col in [5, 6, 8, 9, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22]:
    ws2.cell(row=total_row, column=col).value = f'=SUM({get_column_letter(col)}2:{get_column_letter(col)}43)'
ws2.cell(row=total_row, column=23).value = f'=SUM(W2:W43)'
ws2.cell(row=total_row, column=23).number_format = '$#,##0'
# Avg rates
ws2.cell(row=total_row, column=7).value = f'=IF(E{total_row}=0,"",F{total_row}/E{total_row})'
ws2.cell(row=total_row, column=7).number_format = '0.0%'
ws2.cell(row=total_row, column=10).value = f'=IF(H{total_row}=0,"",I{total_row}/H{total_row})'
ws2.cell(row=total_row, column=10).number_format = '0.0%'

for col in range(1, 25):
    ws2.cell(row=total_row, column=col).fill = total_fill
    ws2.cell(row=total_row, column=col).border = thin_border

style_header(ws2, max_col=24)
style_cells(ws2, end_row=43, max_col=24)

widths2 = [12, 12, 8, 12, 12, 12, 12, 14, 12, 12, 12, 14, 14, 14, 14, 14, 14, 14, 14, 14, 16, 8, 10, 25]
for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

ws2.freeze_panes = 'A2'

# ===================== SHEET 3: WEEKLY DASHBOARD =====================
ws3 = wb.create_sheet("Weekly Dashboard")

headers3 = ["Week #", "Phase", "Total DMs Sent", "Total DM Replies", "DM Reply Rate",
            "Total Emails Sent", "Total Email Opens", "Email Open Rate",
            "Total LinkedIn Outreach", "LinkedIn Accept Rate", "Total LP Visits",
            "Total Sales", "Total Revenue", "Conversion Rate", "Notes"]

for i, h in enumerate(headers3, 1):
    ws3.cell(row=1, column=i, value=h)

for wk in range(1, 7):
    row = wk + 1
    ws3.cell(row=row, column=1, value=wk)
    ws3.cell(row=row, column=2, value=phases.get(wk, "Action"))
    
    start_r = 2 + (wk - 1) * 7
    end_r = start_r + 6
    
    ws3.cell(row=row, column=3).value = f"=SUM('Daily Activity Log'!E{start_r}:E{end_r})"
    ws3.cell(row=row, column=4).value = f"=SUM('Daily Activity Log'!F{start_r}:F{end_r})"
    ws3.cell(row=row, column=5).value = f'=IF(C{row}=0,"",D{row}/C{row})'
    ws3.cell(row=row, column=5).number_format = '0.0%'
    ws3.cell(row=row, column=6).value = f"=SUM('Daily Activity Log'!H{start_r}:H{end_r})"
    ws3.cell(row=row, column=7).value = f"=SUM('Daily Activity Log'!I{start_r}:I{end_r})"
    ws3.cell(row=row, column=8).value = f'=IF(F{row}=0,"",G{row}/F{row})'
    ws3.cell(row=row, column=8).number_format = '0.0%'
    ws3.cell(row=row, column=9).value = f"=SUM('Daily Activity Log'!M{start_r}:M{end_r})"
    ws3.cell(row=row, column=10).value = f"=IF(I{row}=0,\"\",SUM('Daily Activity Log'!N{start_r}:N{end_r})/I{row})"
    ws3.cell(row=row, column=10).number_format = '0.0%'
    ws3.cell(row=row, column=11).value = f"=SUM('Daily Activity Log'!U{start_r}:U{end_r})"
    ws3.cell(row=row, column=12).value = f"=SUM('Daily Activity Log'!V{start_r}:V{end_r})"
    ws3.cell(row=row, column=13).value = f'=L{row}*97'
    ws3.cell(row=row, column=13).number_format = '$#,##0'
    ws3.cell(row=row, column=14).value = f'=IF(K{row}=0,"",L{row}/K{row})'
    ws3.cell(row=row, column=14).number_format = '0.0%'

# Grand total
tr = 8
ws3.cell(row=tr, column=1, value="GRAND TOTAL")
ws3.cell(row=tr, column=1).font = Font(name='Calibri', bold=True, size=11)
for col in [3, 4, 6, 7, 9, 11, 12]:
    ws3.cell(row=tr, column=col).value = f'=SUM({get_column_letter(col)}2:{get_column_letter(col)}7)'
ws3.cell(row=tr, column=13).value = f'=L{tr}*97'
ws3.cell(row=tr, column=13).number_format = '$#,##0'
ws3.cell(row=tr, column=5).value = f'=IF(C{tr}=0,"",D{tr}/C{tr})'
ws3.cell(row=tr, column=5).number_format = '0.0%'
ws3.cell(row=tr, column=8).value = f'=IF(F{tr}=0,"",G{tr}/F{tr})'
ws3.cell(row=tr, column=8).number_format = '0.0%'
ws3.cell(row=tr, column=14).value = f'=IF(K{tr}=0,"",L{tr}/K{tr})'
ws3.cell(row=tr, column=14).number_format = '0.0%'

for col in range(1, 16):
    ws3.cell(row=tr, column=col).fill = total_fill
    ws3.cell(row=tr, column=col).border = thin_border

style_header(ws3, max_col=15)
style_cells(ws3, end_row=7, max_col=15)

widths3 = [10, 12, 14, 14, 12, 14, 14, 12, 16, 14, 14, 10, 12, 12, 25]
for i, w in enumerate(widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

ws3.freeze_panes = 'A2'

# ===================== SHEET 4: EMAIL SEQUENCE TRACKER =====================
ws4 = wb.create_sheet("Email Sequence Tracker")

headers4 = ["Email #", "Email Name", "Phase", "Subject Line A", "Subject Line B",
            "Send Date", "Total Sent", "Opens (A)", "Open Rate (A)", "Opens (B)",
            "Open Rate (B)", "Winner", "Total Clicks", "CTR", "Unsubscribes",
            "Replies", "Sales Attributed", "Notes"]

for i, h in enumerate(headers4, 1):
    ws4.cell(row=1, column=i, value=h)

emails = [
    [1, "The Pain", "Interest", "Are you renting clients or building a business?", "The $48,000 mistake most salon owners make every year"],
    [2, "Social Proof", "Interest", "23 clients came back. Zero ad spend.", "How Fatima stopped paying $4,000/month for ads"],
    [3, "The Tease", "Interest", "The 8-module system that filled 500+ salon calendars", "Before I show you the price \u2014 here's what's inside"],
    [4, "Offer Reveal", "Desire", "Here's everything you get \u2014 and the price drops today", "$1,041 in value. $97 today."],
    [5, "Bonus Stack", "Desire", "Here's what you lose if you wait", "The bonuses disappear before the price goes up"],
    [6, "Testimonials", "Desire", "Real salon owners. Real numbers.", "From $4,000/month in ads to $0 \u2014 their stories"],
    [7, "48hr Warning", "Action", "48 hours left \u2014 then the price goes back to $497", "The math on waiting vs. enrolling right now"],
    [8, "24hr Warning", "Action", "Last chance \u2014 24 hours before the price changes", "What if it doesn't work? (Read this guarantee)"],
    [9, "Final Close", "Action", "This is it \u2014 closing in 6 hours [FINAL NOTICE]", "Final email. Then it's done."],
]

for i, edata in enumerate(emails, 2):
    for j, val in enumerate(edata, 1):
        ws4.cell(row=i, column=j, value=val)
    # Formula columns
    ws4.cell(row=i, column=9).value = f'=IF(G{i}=0,"",H{i}/(G{i}/2))'
    ws4.cell(row=i, column=9).number_format = '0.0%'
    ws4.cell(row=i, column=11).value = f'=IF(G{i}=0,"",J{i}/(G{i}/2))'
    ws4.cell(row=i, column=11).number_format = '0.0%'
    ws4.cell(row=i, column=14).value = f'=IF(G{i}=0,"",M{i}/G{i})'
    ws4.cell(row=i, column=14).number_format = '0.0%'

style_header(ws4, max_col=18)
style_cells(ws4, end_row=10, max_col=18)

widths4 = [8, 14, 10, 35, 35, 12, 10, 10, 10, 10, 10, 8, 12, 8, 12, 10, 14, 25]
for i, w in enumerate(widths4, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w

ws4.freeze_panes = 'A2'

# ===================== SHEET 5: CHANNEL PERFORMANCE =====================
ws5 = wb.create_sheet("Channel Performance")

headers5 = ["Channel", "Total Outreach", "Total Responses", "Response Rate",
            "LP Clicks", "Click Rate", "Sales", "Conversion Rate",
            "Revenue", "Time (hrs)", "Revenue/Hour", "Priority", "Notes"]

for i, h in enumerate(headers5, 1):
    ws5.cell(row=1, column=i, value=h)

channels = ["Instagram DM", "Cold Email (GHL)", "LinkedIn", "Facebook Groups", "WhatsApp", "Threads"]
for i, ch in enumerate(channels, 2):
    ws5.cell(row=i, column=1, value=ch)
    ws5.cell(row=i, column=4).value = f'=IF(B{i}=0,"",C{i}/B{i})'
    ws5.cell(row=i, column=4).number_format = '0.0%'
    ws5.cell(row=i, column=6).value = f'=IF(B{i}=0,"",E{i}/B{i})'
    ws5.cell(row=i, column=6).number_format = '0.0%'
    ws5.cell(row=i, column=8).value = f'=IF(E{i}=0,"",G{i}/E{i})'
    ws5.cell(row=i, column=8).number_format = '0.0%'
    ws5.cell(row=i, column=9).value = f'=G{i}*97'
    ws5.cell(row=i, column=9).number_format = '$#,##0'
    ws5.cell(row=i, column=11).value = f'=IF(J{i}=0,"",I{i}/J{i})'
    ws5.cell(row=i, column=11).number_format = '$#,##0'

# Total row
tr5 = 8
ws5.cell(row=tr5, column=1, value="TOTAL")
ws5.cell(row=tr5, column=1).font = Font(bold=True)
for col in [2, 3, 5, 7, 10]:
    ws5.cell(row=tr5, column=col).value = f'=SUM({get_column_letter(col)}2:{get_column_letter(col)}7)'
ws5.cell(row=tr5, column=9).value = f'=SUM(I2:I7)'
ws5.cell(row=tr5, column=9).number_format = '$#,##0'
for col in range(1, 14):
    ws5.cell(row=tr5, column=col).fill = total_fill
    ws5.cell(row=tr5, column=col).border = thin_border

style_header(ws5, max_col=13)
style_cells(ws5, end_row=7, max_col=13)

widths5 = [18, 14, 14, 12, 10, 10, 8, 12, 10, 10, 12, 8, 25]
for i, w in enumerate(widths5, 1):
    ws5.column_dimensions[get_column_letter(i)].width = w

ws5.freeze_panes = 'A2'

# ===================== SHEET 6: EXECUTION CALENDAR =====================
ws6 = wb.create_sheet("6-Week Execution Calendar")

headers6 = ["Date", "Day", "Week #", "Phase", "Morning (9-11 AM)", "Midday (11 AM-1 PM)",
            "Afternoon (2-4 PM)", "Evening (4-6 PM)", "Key Deliverable", "Status"]

for i, h in enumerate(headers6, 1):
    ws6.cell(row=1, column=i, value=h)

# Task templates by week
tasks = {
    1: {  # Week 1 - Setup + Launch
        "mon": ["Set up GHL sub-account, configure pipeline stages and tags", "Set up Apollo.io, save search filters, start building email list", "Join 8 Facebook groups, optimize Instagram bio + link", "Prepare first 3 Instagram DM script variants", "All platforms set up and configured"],
        "tue": ["Warm email domain (start process), install G Maps Extractor", "Scrape first 2 cities: Miami + LA (Google Maps)", "Start Apollo.io list build (target: 200 contacts today)", "Verify scraped data, clean email list", "200+ leads scraped and verified"],
        "wed": ["Send first batch of Instagram DMs (30 DMs)", "Continue Apollo.io list build (target: 300 more)", "Scrape 2 more cities: Houston + NYC", "Comment on 5+ posts in Facebook groups", "30 DMs sent, 500 leads total"],
        "thu": ["Instagram DMs (40 DMs), track responses", "Load first email batch into GHL (do NOT send yet)", "Scrape 2 more cities: Chicago + Dallas", "Continue Facebook group commenting", "40 DMs sent, 700+ leads"],
        "fri": ["Instagram DMs (50 DMs), review reply rate", "Finalize email list in GHL (500+ contacts)", "Scrape 2 more cities: Atlanta + Tampa", "Review Week 1 progress, adjust scripts if needed", "50 DMs sent, 1000+ leads in database"],
        "sat": ["Light DM follow-ups to replies", "Organize and clean all lead data", "Content planning for next week", "Rest / review metrics", "Database organized"],
        "sun": ["Plan Week 2 priorities", "Review DM reply rate and best-performing opener", "Prep email sequences in GHL for Week 2 launch", "Rest", "Week 2 plan ready"],
    },
    2: {  # Week 2 - Scale Outreach
        "mon": ["Instagram DMs (50/day), A/B test openers", "Launch first cold email batch in GHL (150/day)", "Start LinkedIn outreach (20 requests/day)", "First organic Facebook group post", "Email + DM + LinkedIn all active"],
        "tue": ["Instagram DMs (50/day), track responses", "Monitor email open rates, adjust subject lines", "LinkedIn connection follow-ups", "Facebook group commenting", "All channels running"],
        "wed": ["Instagram DMs (40/day) + reply to responses", "Cold emails (150/day), review A/B results", "LinkedIn messages to new connections", "Second Facebook group post", "Mid-week metric review"],
        "thu": ["Instagram DMs (40/day)", "Cold emails (200/day) — double down on winning subject line", "LinkedIn outreach (25/day)", "Facebook group commenting + engagement", "200+ emails/day milestone"],
        "fri": ["Instagram DMs (40/day) + nurture warm leads", "Review email metrics: opens, clicks, replies", "LinkedIn (20/day)", "WEEKLY REVIEW: DM rate, email rate, LP visits", "Weekly KPI review complete"],
        "sat": ["Light follow-ups only", "Prep Phase 2 (Interest) materials", "Organize warm leads list", "Content review", "Warm leads identified"],
        "sun": ["Plan Phase 2 transition", "Review and segment leads by tier", "Prep nurture sequences", "Rest", "Phase 2 plan ready"],
    },
    3: {  # Week 3 - Nurture
        "mon": ["Launch Email Sequence 1-3 in GHL (nurture phase)", "Instagram DMs (25/day) — nurture warm leads", "Start Instagram 5-day story series (Day 1: RANK)", "LinkedIn posts 2x this week", "Nurture emails live"],
        "tue": ["Instagram Story Day 2: CONVERT", "DMs (25/day) — value-add messages to warm leads", "Monitor email opens/clicks on Email 1", "Facebook group post + engagement", "Story series running"],
        "wed": ["Instagram Story Day 3: REACTIVATE", "DMs (20/day) — follow up with story viewers", "LinkedIn post #1 this week", "Review Email 1 performance", "Mid-week nurture check"],
        "thu": ["Instagram Story Day 4: REFER", "DMs (20/day)", "Monitor Email 2 performance", "Facebook group commenting", "Nurture progressing"],
        "fri": ["Instagram Story Day 5: SYSTEM OVERVIEW", "DMs (20/day)", "WEEKLY REVIEW: warm lead count, email engagement", "Prepare Phase 3 materials if warm leads > 50", "Weekly review + Phase 3 prep"],
        "sat": ["Light follow-ups", "Review story engagement metrics", "Segment leads: who engaged most?", "Rest", "Lead segments updated"],
        "sun": ["Plan Week 4 priorities", "Review warm vs cold lead ratios", "Adjust strategy if warm leads < 50", "Rest", "Week 4 plan ready"],
    },
    4: {  # Week 4 - Deep Nurture
        "mon": ["Continue nurture DMs (20/day) — value content", "Monitor Email 3 performance in GHL", "LinkedIn post + engagement", "Facebook group post", "Nurture continues"],
        "tue": ["DMs (20/day) — check-in messages", "Prep Email 4-6 (Desire phase) in GHL", "LinkedIn connection follow-ups", "Engage with warm leads on their posts", "Desire emails prepped"],
        "wed": ["DMs (20/day) — share success stories", "Review all email metrics so far", "LinkedIn engagement", "Facebook group commenting", "Email metrics reviewed"],
        "thu": ["DMs (15/day) — focus on warm leads only", "CRITICAL: Count warm leads. If < 50, add more cold outreach", "Prep WhatsApp Business account", "Plan live elements for Phase 3", "Warm lead count assessed"],
        "fri": ["DMs (15/day)", "WEEKLY REVIEW: total warm leads, engagement rates", "Finalize Phase 3 launch plan", "Prepare urgency messaging", "Phase 3 launch plan finalized"],
        "sat": ["Light follow-ups", "Organize all warm leads for Phase 3 push", "Test all Phase 3 emails", "Rest", "Phase 3 ready to launch"],
        "sun": ["Final Phase 3 prep", "Review all scripts and messaging", "Mental prep for push week", "Rest", "Ready for Phase 3"],
    },
    5: {  # Week 5 - Desire
        "mon": ["Launch Email 4 (Offer Reveal) in GHL", "Urgency DMs to all warm leads (15-20/day)", "First WhatsApp voice notes to HOT leads", "LinkedIn urgency post", "Offer revealed across all channels"],
        "tue": ["Monitor Email 4 performance", "DMs (15/day) — respond to all replies FAST", "WhatsApp follow-ups (5-10/day)", "Facebook group — share testimonial post", "Responses tracked"],
        "wed": ["Launch Email 5 (Bonus Stack)", "DMs (15/day) — highlight bonuses expiring", "WhatsApp voice notes to new warm leads", "LinkedIn post about results", "Bonus urgency live"],
        "thu": ["Monitor Email 5, prep Email 6", "DMs (10/day) — personal touch messages", "WhatsApp follow-ups", "Facebook group engagement", "Desire building"],
        "fri": ["Launch Email 6 (Testimonials)", "DMs (10/day) — share social proof", "WEEKLY REVIEW: sales count, conversion rate", "Prep Phase 4 (close) materials", "Phase 4 prep started"],
        "sat": ["Count sales so far. If < 10, extend deadline by 5 days", "Final WhatsApp push to undecided leads", "Review what's working, double down", "Rest", "Sales assessment"],
        "sun": ["Plan close week (Phase 4)", "Prep all final emails and DMs", "Organize non-buyer list for re-engagement", "Rest", "Close week ready"],
    },
    6: {  # Week 6 - Close
        "mon": ["Launch Email 7 (48hr Warning)", "DM re-engagement of ALL non-buyers", "WhatsApp final reminders to warm leads", "LinkedIn final post", "48hr countdown started"],
        "tue": ["Launch Email 8 (24hr Warning)", "DMs (10/day) — personal last-chance messages", "WhatsApp text messages with link", "Monitor real-time sales", "24hr countdown"],
        "wed": ["Launch Email 9 (Final Close — 6hrs before)", "Last round of DMs to fence-sitters", "WhatsApp final voice notes", "CLOSE THE WINDOW — remove discount", "CAMPAIGN CLOSED"],
        "thu": ["Send thank-you message to all buyers", "Document final metrics across all channels", "Calculate ROI and channel performance", "Begin post-campaign analysis", "Final metrics documented"],
        "fri": ["Complete campaign post-mortem", "Update Channel Performance sheet", "Identify top-performing channels for next campaign", "Plan next steps / repeat campaign", "Campaign complete"],
        "sat": ["Rest + celebrate results", "", "", "", ""],
        "sun": ["", "", "", "", ""],
    },
}

day_keys = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]
row_num = 2
for wk in range(1, 7):
    for d_idx, dk in enumerate(day_keys):
        d = start_date + timedelta(days=(wk - 1) * 7 + d_idx)
        t = tasks.get(wk, {}).get(dk, ["", "", "", "", ""])
        if len(t) < 5:
            t = t + [""] * (5 - len(t))
        
        ws6.cell(row=row_num, column=1, value=d.strftime("%Y-%m-%d"))
        ws6.cell(row=row_num, column=2, value=days_of_week[d_idx])
        ws6.cell(row=row_num, column=3, value=wk)
        ws6.cell(row=row_num, column=4, value=phases.get(wk, "Action"))
        ws6.cell(row=row_num, column=5, value=t[0])
        ws6.cell(row=row_num, column=6, value=t[1])
        ws6.cell(row=row_num, column=7, value=t[2])
        ws6.cell(row=row_num, column=8, value=t[3])
        ws6.cell(row=row_num, column=9, value=t[4])
        ws6.cell(row=row_num, column=10, value="Not Started")
        row_num += 1

# Data validation for status
dv_status6 = DataValidation(type="list", formula1='"Not Started,In Progress,Done"')
ws6.add_data_validation(dv_status6)
dv_status6.add(f"J2:J43")

style_header(ws6, max_col=10)
style_cells(ws6, end_row=43, max_col=10)

widths6 = [12, 12, 8, 12, 40, 40, 40, 35, 30, 12]
for i, w in enumerate(widths6, 1):
    ws6.column_dimensions[get_column_letter(i)].width = w

ws6.freeze_panes = 'A2'

# SAVE
path = "/Users/nishamore/Documents/Salon world wide/Outreach_Pipeline_Tracker.xlsx"
wb.save(path)
print(f"Outreach Pipeline Tracker saved to: {path}")
print(f"File size: {os.path.getsize(path)} bytes")
